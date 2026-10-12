import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path="DataRequestStudio_Executive_Presentation.pptx"):
    prs = Presentation()
    # 16:9 Widescreen standard
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # McKinsey-Inspired Color Palette
    NAVY_DARK = RGBColor(15, 32, 66)      # #0F2042 - Primary McKinsey Navy
    NAVY_MID = RGBColor(26, 54, 93)       # #1A365D - Secondary Navy
    MCK_BLUE = RGBColor(0, 119, 200)      # #0077C8 - McKinsey Accent Blue
    CYAN_ACCENT = RGBColor(2, 132, 199)   # #0284C7 - Bright Blue Highlight
    CHARCOAL = RGBColor(30, 41, 59)       # #1E293B - Deep Body Text
    SLATE = RGBColor(71, 85, 105)         # #475569 - Secondary Muted Text
    LIGHT_BG = RGBColor(248, 250, 252)    # #F8FAFC - Card Background
    CARD_BORDER = RGBColor(226, 232, 240) # #E2E8F0 - Border Stroke
    WHITE = RGBColor(255, 255, 255)       # #FFFFFF - Pure White
    GREEN_ACC = RGBColor(16, 149, 119)    # #109577 - Success / Safe
    AMBER_ACC = RGBColor(217, 119, 6)     # #D97706 - Warning / Attention

    FONT_FAMILY = "Arial"

    def add_header(slide, category_text, action_title_text, subtitle_text=""):
        # Category Tracker
        tb_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.733), Inches(0.3))
        tf_cat = tb_cat.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_right = tf_cat.margin_top = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.name = FONT_FAMILY
        p_cat.font.size = Pt(9.5)
        p_cat.font.bold = True
        p_cat.font.color.rgb = MCK_BLUE

        # Action Title
        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.733), Inches(0.6))
        tf_title = tb_title.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = action_title_text
        p_title.font.name = FONT_FAMILY
        p_title.font.size = Pt(19)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY_DARK

        # Subtitle / context if available
        if subtitle_text:
            tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.35))
            tf_sub = tb_sub.text_frame
            tf_sub.word_wrap = True
            tf_sub.margin_left = tf_sub.margin_right = tf_sub.margin_top = tf_sub.margin_bottom = 0
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = subtitle_text
            p_sub.font.name = FONT_FAMILY
            p_sub.font.size = Pt(11)
            p_sub.font.color.rgb = SLATE

        # Subtle separator line
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.8), Inches(1.72 if subtitle_text else 1.45), Inches(11.733), Inches(0.015)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = CARD_BORDER
        line.line.color.rgb = CARD_BORDER

    def add_footer(slide, current_slide, total_slides=10):
        # Footer text left
        tb_l = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(6.0), Inches(0.25))
        tf_l = tb_l.text_frame
        tf_l.margin_left = tf_l.margin_right = tf_l.margin_top = tf_l.margin_bottom = 0
        p_l = tf_l.paragraphs[0]
        p_l.text = "DataRequestStudio | Transformasi Tata Kelola Data CEISA Data Warehouse"
        p_l.font.name = FONT_FAMILY
        p_l.font.size = Pt(8.5)
        p_l.font.color.rgb = SLATE

        # Footer text center
        tb_c = slide.shapes.add_textbox(Inches(5.0), Inches(7.05), Inches(3.333), Inches(0.25))
        tf_c = tb_c.text_frame
        tf_c.margin_left = tf_c.margin_right = tf_c.margin_top = tf_c.margin_bottom = 0
        p_c = tf_c.paragraphs[0]
        p_c.alignment = PP_ALIGN.CENTER
        p_c.text = "Klasifikasi: Terbatas / Khusus DJBC & Kemenkeu"
        p_c.font.name = FONT_FAMILY
        p_c.font.size = Pt(8.5)
        p_c.font.color.rgb = SLATE

        # Footer text right (slide number)
        tb_r = slide.shapes.add_textbox(Inches(10.533), Inches(7.05), Inches(2.0), Inches(0.25))
        tf_r = tb_r.text_frame
        tf_r.margin_left = tf_r.margin_right = tf_r.margin_top = tf_r.margin_bottom = 0
        p_r = tf_r.paragraphs[0]
        p_r.alignment = PP_ALIGN.RIGHT
        p_r.text = f"Slide {current_slide} dari {total_slides}"
        p_r.font.name = FONT_FAMILY
        p_r.font.size = Pt(8.5)
        p_r.font.color.rgb = SLATE

    # =========================================================================
    # SLIDE 1: TITLE SLIDE (Executive Dark McKinsey Theme)
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)

    # Dark background fill
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY_DARK
    bg1.line.color.rgb = NAVY_DARK

    # Blue McKinsey accent bar on top left
    accent_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.3), Inches(1.2), Inches(0.08))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = MCK_BLUE
    accent_bar.line.color.rgb = MCK_BLUE

    # Supertitle
    tb_super = slide1.shapes.add_textbox(Inches(0.8), Inches(1.55), Inches(11.5), Inches(0.4))
    tf_super = tb_super.text_frame
    tf_super.word_wrap = True
    p_sup = tf_super.paragraphs[0]
    p_sup.text = "DIREKTORAT JENDERAL BEA DAN CUKAI  |  ENTERPRISE DATA GOVERNANCE"
    p_sup.font.name = FONT_FAMILY
    p_sup.font.size = Pt(11)
    p_sup.font.bold = True
    p_sup.font.color.rgb = MCK_BLUE

    # Main Title
    tb_main = slide1.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.5), Inches(1.8))
    tf_main = tb_main.text_frame
    tf_main.word_wrap = True
    p_main = tf_main.paragraphs[0]
    p_main.text = "DataRequestStudio:\nTata Kelola & Otomasi Pemenuhan Data"
    p_main.font.name = FONT_FAMILY
    p_main.font.size = Pt(36)
    p_main.font.bold = True
    p_main.font.color.rgb = WHITE

    # Subtitle
    tb_desc = slide1.shapes.add_textbox(Inches(0.8), Inches(3.9), Inches(11.5), Inches(0.9))
    tf_desc = tb_desc.text_frame
    tf_desc.word_wrap = True
    p_desc = tf_desc.paragraphs[0]
    p_desc.text = "Kompendium Eksekutif: Definisi (What), Urgensi & Nilai Tambah (Why), serta Panduan Akses (How to Access) Portal Permohonan Data CEISA Data Warehouse"
    p_desc.font.name = FONT_FAMILY
    p_desc.font.size = Pt(14)
    p_desc.font.color.rgb = RGBColor(203, 213, 225)

    # 3 Metadata Badges / Info Cards
    meta_cards = [
        ("FOKUS STRATEGIS", "Otomasi SLA & Kepatuhan UU PDP No. 27/2022"),
        ("INTEGRASI SISTEM", "CEISA Data Warehouse & OpenMetadata Enterprise"),
        ("STATUS DOKUMEN", "Versi 2.0 Enterprise  |  Oktober 2026")
    ]
    for i, (m_title, m_val) in enumerate(meta_cards):
        c_left = Inches(0.8 + i * 4.0)
        c_top = Inches(5.3)
        c_box = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_left, c_top, Inches(3.7), Inches(1.1))
        c_box.fill.solid()
        c_box.fill.fore_color.rgb = RGBColor(26, 46, 85)
        c_box.line.color.rgb = RGBColor(40, 68, 115)
        c_box.line.width = Pt(1)

        tb_m = slide1.shapes.add_textbox(c_left + Inches(0.2), c_top + Inches(0.15), Inches(3.3), Inches(0.8))
        tf_m = tb_m.text_frame
        tf_m.word_wrap = True
        p_mt = tf_m.paragraphs[0]
        p_mt.text = m_title
        p_mt.font.name = FONT_FAMILY
        p_mt.font.size = Pt(9)
        p_mt.font.bold = True
        p_mt.font.color.rgb = MCK_BLUE

        p_mv = tf_m.add_paragraph()
        p_mv.text = m_val
        p_mv.font.name = FONT_FAMILY
        p_mv.font.size = Pt(11)
        p_mv.font.color.rgb = WHITE
        p_mv.space_before = Pt(4)

    # =========================================================================
    # SLIDE 2: EXECUTIVE SUMMARY (Pyramid Principle)
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(
        slide2,
        "RINGKASAN EKSEKUTIF",
        "DataRequestStudio mentransformasi pemenuhan data DJBC dari birokrasi manual menjadi platform swalayan yang aman, patuh PDP, dan terotomasi.",
        "Pendekatan 'Pyramid Principle': Inti proposisi nilai terbagi dalam 3 pilar operasional utama."
    )

    pillars = [
        (
            "WHAT: PLATFORM PORTAL",
            "Portal Terpadu & Terintegrasi",
            [
                "Portal otomasi end-to-end permintaan data kepabeanan & cukai.",
                "Terintegrasi langsung dengan Katalog Enterprise OpenMetadata dan CEISA Data Warehouse.",
                "Engine query multi-dialek (PostgreSQL, Oracle, SQLite) dengan validasi skema otomatis.",
                "Sistem repositori (CMS) pemantauan permohonan dan repositori data aset dinas."
            ],
            NAVY_MID
        ),
        (
            "WHY: URGENSI & BENEFIT",
            "Reduksi Lead Time & Jaminan PDP",
            [
                "Memangkas waktu tunggu dari 14-30 hari kerja menjadi <24-48 jam.",
                "Penegakan otomatis UU PDP No. 27/2022 melalui Dynamic Field Masking (NIK, NPWP, Rekening).",
                "Menghilangkan penumpukan tiket ad-hoc dan beban kueri manual tim DWH.",
                "100% audit trail transparan dengan tanda tangan digital & verifikasi berjenjang."
            ],
            MCK_BLUE
        ),
        (
            "HOW TO ACCESS: ALUR AKSES",
            "Akses Swalayan Berbasis Peran",
            [
                "Akses langsung melalui portal peramban web internal DJBC / demo publik.",
                "Otentikasi aman berbasis Single Sign-On (SSO) Kemenkeu & RBAC 4-peran.",
                "Alur 4 tahap: Eksplorasi katalog -> Penyusunan form -> Review Steward -> Unduh aman.",
                "Dukungan helpdesk terpadu Subdirektorat DWH dengan SLA respon 2 jam."
            ],
            NAVY_DARK
        )
    ]

    for idx, (p_tag, p_head, p_points, header_color) in enumerate(pillars):
        c_left = Inches(0.8 + idx * 4.0)
        c_top = Inches(1.9)
        c_w = Inches(3.7)
        c_h = Inches(4.0)

        # Card container
        card = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_left, c_top, c_w, c_h)
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        # Header bar of card
        c_bar = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_left, c_top, c_w, Inches(0.75))
        c_bar.fill.solid()
        c_bar.fill.fore_color.rgb = header_color
        c_bar.line.color.rgb = header_color

        tb_ch = slide2.shapes.add_textbox(c_left + Inches(0.2), c_top + Inches(0.08), c_w - Inches(0.4), Inches(0.6))
        tf_ch = tb_ch.text_frame
        tf_ch.word_wrap = True
        p_ch1 = tf_ch.paragraphs[0]
        p_ch1.text = p_tag
        p_ch1.font.name = FONT_FAMILY
        p_ch1.font.size = Pt(8.5)
        p_ch1.font.bold = True
        p_ch1.font.color.rgb = RGBColor(203, 213, 225)

        p_ch2 = tf_ch.add_paragraph()
        p_ch2.text = p_head
        p_ch2.font.name = FONT_FAMILY
        p_ch2.font.size = Pt(11)
        p_ch2.font.bold = True
        p_ch2.font.color.rgb = WHITE

        # Body points
        tb_cb = slide2.shapes.add_textbox(c_left + Inches(0.2), c_top + Inches(0.9), c_w - Inches(0.4), Inches(3.0))
        tf_cb = tb_cb.text_frame
        tf_cb.word_wrap = True
        tf_cb.margin_left = tf_cb.margin_right = tf_cb.margin_top = tf_cb.margin_bottom = 0

        for b_idx, bullet in enumerate(p_points):
            p_b = tf_cb.paragraphs[0] if b_idx == 0 else tf_cb.add_paragraph()
            p_b.text = f"•  {bullet}"
            p_b.font.name = FONT_FAMILY
            p_b.font.size = Pt(10)
            p_b.font.color.rgb = CHARCOAL
            p_b.space_after = Pt(8)

    # Bottom Callout Banner
    bot_banner = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.1), Inches(11.733), Inches(0.7))
    bot_banner.fill.solid()
    bot_banner.fill.fore_color.rgb = RGBColor(238, 242, 255)
    bot_banner.line.color.rgb = RGBColor(199, 210, 254)
    bot_banner.line.width = Pt(1)

    tb_bb = slide2.shapes.add_textbox(Inches(1.0), Inches(6.15), Inches(11.333), Inches(0.6))
    tf_bb = tb_bb.text_frame
    tf_bb.word_wrap = True
    p_bb = tf_bb.paragraphs[0]
    p_bb.text = "KESIMPULAN EKSEKUTIF: DataRequestStudio menjawab kesenjangan antara lonjakan kebutuhan analitik dinas dan tuntutan kepatuhan data privat, menghasilkan efisiensi tata kelola data hingga 85%."
    p_bb.font.name = FONT_FAMILY
    p_bb.font.size = Pt(10.5)
    p_bb.font.bold = True
    p_bb.font.color.rgb = NAVY_MID

    add_footer(slide2, 2)

    # =========================================================================
    # SLIDE 3: WHAT IS DATAREQUESTSTUDIO? (Arsitektur & 4 Modul)
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(
        slide3,
        "WHAT — DEFINISI & ARSITEKTUR PLATFORM",
        "Platform mengintegrasikan katalog semantik DWH, engine kepatuhan PDP, dan repositori permohonan ke dalam satu sistem terpadu.",
        "Arsitektur modular menjamin keandalan skalabilitas, interoperabilitas multi-database, dan integritas data."
    )

    modules = [
        (
            "1. Enterprise Semantic Catalog",
            "Integrasi OpenMetadata & CEISA DWH",
            "Menyediakan navigasi visual terhadap skema database kepabeanan dan cukai. Pengguna dapat mengeksplorasi kamus data (data dictionary), lineage kolom, tipe data, serta deskripsi bisnis tanpa perlu mengakses server basis data secara langsung.",
            MCK_BLUE
        ),
        (
            "2. Multi-Role Governance Engine",
            "Pemisahan Tugas & Persetujuan Berjenjang",
            "Menerapkan prinsip 'Separation of Duties' dengan 4 peran spesifik: Requester, Data Steward, Data Analyst, dan Data Engineer. Setiap alur permohonan melalui verifikasi justifikasi dinas, review kepatuhan, dan penandatanganan digital resmi.",
            NAVY_DARK
        ),
        (
            "3. Automated PDP & Security Masking",
            "Proteksi Otomatis Data Pribadi",
            "Mendeteksi secara proaktif atribut bernilai pribadi (PII) seperti NIK, NPWP pribadi, nomor paspor, dan rekening bank. Sistem menerapkan masking dinamis dan tokenisasi sebelum hasil diekstrak ke dalam format Excel atau ZIP terenkripsi.",
            GREEN_ACC
        ),
        (
            "4. CMS Repository & Audit Trail",
            "Pusat Pemantauan & Jejak Rekam Digital",
            "Repositori terpusat untuk memantau status permohonan (Disetujui, Perlu Revisi, Ditolak). Dilengkapi metrik latensi DWH, SLA counter, pencarian instan, dan log audit forensik lengkap yang dapat dipertanggungjawabkan kepada auditor.",
            AMBER_ACC
        )
    ]

    for idx, (m_title, m_sub, m_desc, m_accent) in enumerate(modules):
        row = idx // 2
        col = idx % 2
        m_left = Inches(0.8 + col * 6.0)
        m_top = Inches(1.9 + row * 2.45)
        m_w = Inches(5.7)
        m_h = Inches(2.25)

        # Card container
        m_card = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, m_left, m_top, m_w, m_h)
        m_card.fill.solid()
        m_card.fill.fore_color.rgb = LIGHT_BG
        m_card.line.color.rgb = CARD_BORDER
        m_card.line.width = Pt(1)

        # Left accent stripe
        m_stripe = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, m_left, m_top, Inches(0.12), m_h)
        m_stripe.fill.solid()
        m_stripe.fill.fore_color.rgb = m_accent
        m_stripe.line.color.rgb = m_accent

        # Text inside card
        tb_mc = slide3.shapes.add_textbox(m_left + Inches(0.3), m_top + Inches(0.15), m_w - Inches(0.5), m_h - Inches(0.3))
        tf_mc = tb_mc.text_frame
        tf_mc.word_wrap = True
        tf_mc.margin_left = tf_mc.margin_right = tf_mc.margin_top = tf_mc.margin_bottom = 0

        p_m1 = tf_mc.paragraphs[0]
        p_m1.text = m_title
        p_m1.font.name = FONT_FAMILY
        p_m1.font.size = Pt(12.5)
        p_m1.font.bold = True
        p_m1.font.color.rgb = NAVY_DARK

        p_m2 = tf_mc.add_paragraph()
        p_m2.text = m_sub
        p_m2.font.name = FONT_FAMILY
        p_m2.font.size = Pt(10)
        p_m2.font.bold = True
        p_m2.font.color.rgb = m_accent
        p_m2.space_after = Pt(6)

        p_m3 = tf_mc.add_paragraph()
        p_m3.text = m_desc
        p_m3.font.name = FONT_FAMILY
        p_m3.font.size = Pt(9.5)
        p_m3.font.color.rgb = CHARCOAL

    add_footer(slide3, 3)

    # =========================================================================
    # SLIDE 4: WHAT - CAKUPAN 6 PILAR PROSES BISNIS
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(
        slide4,
        "WHAT — CAKUPAN PROSES BISNIS & DATA DOMAIN",
        "DataRequestStudio mencakup seluruh pilar proses bisnis kepabeanan dan cukai dengan pemetaan tabel terstandarisasi di CEISA DWH.",
        "Klasifikasi entitas memudahkan pemohon memilih parameter data sesuai yurisdiksi dan kewenangan tugas."
    )

    biz_domains = [
        (
            "1. IMPOR",
            "PIB BC 2.0 & Dokumen Pelengkap",
            [
                "Header & Dokumen PIB BC 2.0",
                "Item Barang, HS Code, & Tarif Bea Masuk",
                "Data Kontainer, Kemasan, & Nilai Pabean",
                "Status Penjaluran (Merah, Kuning, Hijau)"
            ]
        ),
        (
            "2. EKSPOR",
            "PEB BC 3.0 & Pengawasan Fisik",
            [
                "Dokumen Ekspor PEB BC 3.0",
                "Laporan Pemeriksaan Surveyor (LPS)",
                "Data Devisa Hasil Ekspor (DHE)",
                "Monitoring Muat & Keberangkatan Sarana"
            ]
        ),
        (
            "3. CUKAI",
            "Barang Kena Cukai (BKC) & Pita Cukai",
            [
                "Dokumen Pemesanan Pita Cukai (CK-1)",
                "Pelunasan Cukai Hasil Tembakau (HT)",
                "Monitoring Pabrik MMEA & Etil Alkohol",
                "Produksi, Tarif, & Pelunasan Cukai"
            ]
        ),
        (
            "4. FASILITAS",
            "Tempat Penimbunan Berikat & KITE",
            [
                "Kawasan Berikat (KB) & Gudang Berikat",
                "Kemudahan Impor Tujuan Ekspor (KITE)",
                "Fasilitas Pembebasan / Keringanan Bea",
                "Monitoring Mutasi Bahan Baku & Sisa"
            ]
        ),
        (
            "5. MANIFEST",
            "Pengawasan Sarana Pengangkut (BC 1.1)",
            [
                "Manifest Kedatangan (Inward Manifest BC 1.1)",
                "Manifest Keberangkatan (Outward Manifest)",
                "Data Sarana Pengangkut (Vessel / Flight)",
                "Bill of Lading (B/L) & Air Waybill (AWB)"
            ]
        ),
        (
            "6. PERBENDAHARAAN",
            "Penerimaan Negara & Piutang Pabean",
            [
                "Billing Penerimaan Bea Masuk, Keluar, Cukai",
                "Piutang, Tagihan Surat Penetapan (SPTNP)",
                "Pengelolaan Jaminan Kepabeanan & Cukai",
                "Restitusi, Pengembalian, & Keberatan"
            ]
        )
    ]

    for idx, (b_name, b_sub, b_items) in enumerate(biz_domains):
        row = idx // 3
        col = idx % 3
        b_left = Inches(0.8 + col * 4.0)
        b_top = Inches(1.9 + row * 2.25)
        b_w = Inches(3.7)
        b_h = Inches(2.05)

        b_card = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, b_left, b_top, b_w, b_h)
        b_card.fill.solid()
        b_card.fill.fore_color.rgb = LIGHT_BG
        b_card.line.color.rgb = CARD_BORDER
        b_card.line.width = Pt(1)

        # Header tag
        tb_bt = slide4.shapes.add_textbox(b_left + Inches(0.2), b_top + Inches(0.12), b_w - Inches(0.4), Inches(0.55))
        tf_bt = tb_bt.text_frame
        tf_bt.word_wrap = True
        tf_bt.margin_left = tf_bt.margin_right = tf_bt.margin_top = tf_bt.margin_bottom = 0

        p_bt1 = tf_bt.paragraphs[0]
        p_bt1.text = b_name
        p_bt1.font.name = FONT_FAMILY
        p_bt1.font.size = Pt(11)
        p_bt1.font.bold = True
        p_bt1.font.color.rgb = NAVY_DARK

        p_bt2 = tf_bt.add_paragraph()
        p_bt2.text = b_sub
        p_bt2.font.name = FONT_FAMILY
        p_bt2.font.size = Pt(9)
        p_bt2.font.bold = True
        p_bt2.font.color.rgb = MCK_BLUE

        # Items
        tb_bi = slide4.shapes.add_textbox(b_left + Inches(0.2), b_top + Inches(0.7), b_w - Inches(0.4), Inches(1.25))
        tf_bi = tb_bi.text_frame
        tf_bi.word_wrap = True
        tf_bi.margin_left = tf_bi.margin_right = tf_bi.margin_top = tf_bi.margin_bottom = 0

        for it_idx, item in enumerate(b_items):
            p_it = tf_bi.paragraphs[0] if it_idx == 0 else tf_bi.add_paragraph()
            p_it.text = f"• {item}"
            p_it.font.name = FONT_FAMILY
            p_it.font.size = Pt(9)
            p_it.font.color.rgb = CHARCOAL
            p_it.space_after = Pt(2)

    # Bottom summary box
    bot_info = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.45), Inches(11.733), Inches(0.45))
    bot_info.fill.solid()
    bot_info.fill.fore_color.rgb = RGBColor(241, 245, 249)
    bot_info.line.color.rgb = CARD_BORDER

    tb_binfo = slide4.shapes.add_textbox(Inches(1.0), Inches(6.5), Inches(11.333), Inches(0.35))
    tf_binfo = tb_binfo.text_frame
    tf_binfo.word_wrap = True
    p_binfo = tf_binfo.paragraphs[0]
    p_binfo.text = "Seluruh 6 domain terpetakan secara otomatis dengan relasi Foreign Key teruji di CEISA DWH, menjamin keakuratan sintaks SQL."
    p_binfo.font.name = FONT_FAMILY
    p_binfo.font.size = Pt(9.5)
    p_binfo.font.color.rgb = SLATE

    add_footer(slide4, 4)

    # =========================================================================
    # SLIDE 5: WHY DATAREQUESTSTUDIO? (Komparasi Sebelum vs Sesudah)
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(
        slide5,
        "WHY — URGENSI & TRANSFORMASI OPERASIONAL",
        "Modernisasi pemenuhan data mengatasi tiga hambatan kritis: bottleneck birokrasi, risiko kebocoran data pribadi, dan fragmentasi kueri.",
        "Perbandingan matriks operasional antara proses konvensional dan ekosistem DataRequestStudio."
    )

    dimensions = [
        (
            "DIMENSI 1: KECEPATAN & SLA PEMENUHAN",
            "Sebelumnya: Proses Manual Berbelit",
            "• Pemohon menyusun Nota Dinas fisik dan menunggu paraf berjenjang antar eselon.\n• Waktu pemenuhan data rata-rata memakan 14 s.d. 30 hari kerja.\n• Beban antrean tinggi di tim pengelola basis data.",
            "Dengan DataRequestStudio: Terotomasi & Mandiri",
            "• Pemohon menyusun dan memvalidasi kueri secara swalayan dalam hitungan menit.\n• Lead time terpangkas drastis menjadi <24 s.d. 48 jam kerja.\n• Verifikasi digital instan mempercepat siklus persetujuan.",
            GREEN_ACC
        ),
        (
            "DIMENSI 2: KEPATUHAN UU PDP & PRIVASI",
            "Sebelumnya: Risiko Paparan PII Tinggi",
            "• Data mentah diekstrak tanpa filter perlindungan data pribadi terstandar.\n• NIK, NPWP pribadi, dan rekening bank berisiko terekspos ke pihak tidak berwenang.\n• Potensi pelanggaran hukum terhadap UU PDP No. 27/2022.",
            "Dengan DataRequestStudio: Dynamic PDP Masking",
            "• Sistem secara otomatis mendeteksi kolom PII dan menerapkan masking ketat.\n• Hanya atribut bisnis yang diizinkan yang diekstrak ke dalam dataset.\n• Jaminan kepatuhan 100% terhadap regulasi perlindungan data negara.",
            MCK_BLUE
        ),
        (
            "DIMENSI 3: AKUNTABILITAS & AUDIT TRAIL",
            "Sebelumnya: Silo & 'Black Box'",
            "• Riwayat permohonan tercecer di email atau arsip nota dinas fisik.\n• Sulit melacak siapa yang meminta data, untuk tujuan apa, dan siapa verifikatornya.\n• Duplikasi permohonan data yang sama sering terjadi berulang kali.",
            "Dengan DataRequestStudio: 100% Transparan & Reusable",
            "• Repositori CMS terpusat mencatat seluruh riwayat, status, dan timestamp.\n• Dilengkapi digital signature verifikator dan justifikasi dinas tercatat permanen.\n• Kueri terstandarisasi dapat digunakan kembali (reusable assets).",
            NAVY_MID
        )
    ]

    for idx, (dim_title, bef_title, bef_desc, aft_title, aft_desc, aft_color) in enumerate(dimensions):
        row_top = Inches(1.9 + idx * 1.5)
        row_h = Inches(1.35)

        # Dimension Label Bar (Left)
        d_lbl = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), row_top, Inches(2.6), row_h)
        d_lbl.fill.solid()
        d_lbl.fill.fore_color.rgb = NAVY_DARK
        d_lbl.line.color.rgb = NAVY_DARK

        tb_dl = slide5.shapes.add_textbox(Inches(0.9), row_top + Inches(0.2), Inches(2.4), row_h - Inches(0.4))
        tf_dl = tb_dl.text_frame
        tf_dl.word_wrap = True
        tf_dl.margin_left = tf_dl.margin_right = tf_dl.margin_top = tf_dl.margin_bottom = 0
        p_dl = tf_dl.paragraphs[0]
        p_dl.text = dim_title
        p_dl.font.name = FONT_FAMILY
        p_dl.font.size = Pt(10)
        p_dl.font.bold = True
        p_dl.font.color.rgb = WHITE

        # Before Card (Middle)
        b_box = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3.55), row_top, Inches(4.3), row_h)
        b_box.fill.solid()
        b_box.fill.fore_color.rgb = RGBColor(254, 242, 242)
        b_box.line.color.rgb = RGBColor(254, 202, 202)
        b_box.line.width = Pt(1)

        tb_bc = slide5.shapes.add_textbox(Inches(3.7), row_top + Inches(0.1), Inches(4.0), row_h - Inches(0.2))
        tf_bc = tb_bc.text_frame
        tf_bc.word_wrap = True
        tf_bc.margin_left = tf_bc.margin_right = tf_bc.margin_top = tf_bc.margin_bottom = 0
        p_bch = tf_bc.paragraphs[0]
        p_bch.text = bef_title
        p_bch.font.name = FONT_FAMILY
        p_bch.font.size = Pt(9.5)
        p_bch.font.bold = True
        p_bch.font.color.rgb = RGBColor(185, 28, 28)

        p_bcb = tf_bc.add_paragraph()
        p_bcb.text = bef_desc
        p_bcb.font.name = FONT_FAMILY
        p_bcb.font.size = Pt(8.5)
        p_bcb.font.color.rgb = CHARCOAL
        p_bcb.space_before = Pt(2)

        # After Card (Right)
        a_box = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.0), row_top, Inches(4.533), row_h)
        a_box.fill.solid()
        a_box.fill.fore_color.rgb = RGBColor(240, 253, 244)
        a_box.line.color.rgb = RGBColor(187, 247, 208)
        a_box.line.width = Pt(1)

        tb_ac = slide5.shapes.add_textbox(Inches(8.15), row_top + Inches(0.1), Inches(4.2), row_h - Inches(0.2))
        tf_ac = tb_ac.text_frame
        tf_ac.word_wrap = True
        tf_ac.margin_left = tf_ac.margin_right = tf_ac.margin_top = tf_ac.margin_bottom = 0
        p_ach = tf_ac.paragraphs[0]
        p_ach.text = aft_title
        p_ach.font.name = FONT_FAMILY
        p_ach.font.size = Pt(9.5)
        p_ach.font.bold = True
        p_ach.font.color.rgb = RGBColor(21, 128, 61)

        p_acb = tf_ac.add_paragraph()
        p_acb.text = aft_desc
        p_acb.font.name = FONT_FAMILY
        p_acb.font.size = Pt(8.5)
        p_acb.font.color.rgb = CHARCOAL
        p_acb.space_before = Pt(2)

    # Bottom Callout Summary
    bot_call = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.5), Inches(11.733), Inches(0.4))
    bot_call.fill.solid()
    bot_call.fill.fore_color.rgb = RGBColor(238, 242, 255)
    bot_call.line.color.rgb = RGBColor(199, 210, 254)

    tb_bcall = slide5.shapes.add_textbox(Inches(1.0), Inches(6.52), Inches(11.333), Inches(0.3))
    tf_bcall = tb_bcall.text_frame
    tf_bcall.word_wrap = True
    p_bcall = tf_bcall.paragraphs[0]
    p_bcall.text = "Transformasi Paradigma: Dari 'Menunggu Tim IT Mengekstrak Data' menjadi 'Tata Kelola Mandiri yang Diawasi Sistem'."
    p_bcall.font.name = FONT_FAMILY
    p_bcall.font.size = Pt(9.5)
    p_bcall.font.bold = True
    p_bcall.font.color.rgb = MCK_BLUE

    add_footer(slide5, 5)

    # =========================================================================
    # SLIDE 6: WHY - DAMPAK STRATEGIS & VALUE DRIVERS
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(
        slide6,
        "WHY — PENGUNGKIT NILAI STRATEGIS (VALUE DRIVERS)",
        "Empat pengungkit nilai utama mendorong efisiensi operasional dan memperkuat integritas tata kelola data DJBC.",
        "Kombinasi efisiensi waktu, mitigasi risiko hukum, dan optimalisasi infrastruktur menghasilkan dampak nyata bagi institusi."
    )

    drivers = [
        (
            "1. EFISIENSI WAKTU & BIAYA",
            "Reduksi 85% Lead Time",
            "Memangkas ratusan jam kerja yang sebelumnya terbuang pada birokrasi manual dan drafting kueri ad-hoc berulang. Tim analis dapat langsung fokus pada perumusan rekomendasi kebijakan.",
            "85%",
            "Penghematan Waktu Siklus",
            MCK_BLUE
        ),
        (
            "2. MITIGASI RISIKO HUKUM",
            "Kepatuhan 100% UU PDP",
            "Penerapan automated data classification dan masking dinamis mengeliminasi risiko sanksi administratif dan hukum terkait kebocoran data rahasia wajib pajak dan pengguna jasa.",
            "100%",
            "Patuh UU No. 27/2022",
            GREEN_ACC
        ),
        (
            "3. KUALITAS & KONSISTENSI",
            "Single Source of Truth",
            "Menghilangkan inkonsistensi data antar unit kerja. Seluruh permohonan mengacu pada definisi skema dan relasi data mart resmi yang tersertifikasi di OpenMetadata CEISA DWH.",
            "0%",
            "Inkonsistensi Logika Data",
            NAVY_MID
        ),
        (
            "4. INTEGRITAS TATA KELOLA",
            "Digital Audit Forensik",
            "Setiap akses tercatat lengkap dengan identitas pemohon, unit kerja, alasan kedinasan, waktu penandatanganan, dan durasi eksekusi kueri, siap diaudit BPK maupun Itjen Kemenkeu.",
            "24/7",
            "Audit Trail Real-Time",
            AMBER_ACC
        )
    ]

    for idx, (d_tag, d_head, d_body, d_stat, d_stat_lbl, d_color) in enumerate(drivers):
        c_left = Inches(0.8 + idx * 3.0)
        c_top = Inches(1.9)
        c_w = Inches(2.8)
        c_h = Inches(4.4)

        card = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_left, c_top, c_w, c_h)
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        # Top metric banner
        stat_box = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_left, c_top, c_w, Inches(1.1))
        stat_box.fill.solid()
        stat_box.fill.fore_color.rgb = d_color
        stat_box.line.color.rgb = d_color

        tb_st = slide6.shapes.add_textbox(c_left, c_top + Inches(0.1), c_w, Inches(0.9))
        tf_st = tb_st.text_frame
        tf_st.word_wrap = True
        p_st1 = tf_st.paragraphs[0]
        p_st1.alignment = PP_ALIGN.CENTER
        p_st1.text = d_stat
        p_st1.font.name = FONT_FAMILY
        p_st1.font.size = Pt(26)
        p_st1.font.bold = True
        p_st1.font.color.rgb = WHITE

        p_st2 = tf_st.add_paragraph()
        p_st2.alignment = PP_ALIGN.CENTER
        p_st2.text = d_stat_lbl
        p_st2.font.name = FONT_FAMILY
        p_st2.font.size = Pt(8.5)
        p_st2.font.color.rgb = RGBColor(226, 232, 240)

        # Body text
        tb_db = slide6.shapes.add_textbox(c_left + Inches(0.18), c_top + Inches(1.25), c_w - Inches(0.36), Inches(3.0))
        tf_db = tb_db.text_frame
        tf_db.word_wrap = True
        tf_db.margin_left = tf_db.margin_right = tf_db.margin_top = tf_db.margin_bottom = 0

        p_dh1 = tf_db.paragraphs[0]
        p_dh1.text = d_tag
        p_dh1.font.name = FONT_FAMILY
        p_dh1.font.size = Pt(8.5)
        p_dh1.font.bold = True
        p_dh1.font.color.rgb = SLATE

        p_dh2 = tf_db.add_paragraph()
        p_dh2.text = d_head
        p_dh2.font.name = FONT_FAMILY
        p_dh2.font.size = Pt(11)
        p_dh2.font.bold = True
        p_dh2.font.color.rgb = NAVY_DARK
        p_dh2.space_after = Pt(6)

        p_dh3 = tf_db.add_paragraph()
        p_dh3.text = d_body
        p_dh3.font.name = FONT_FAMILY
        p_dh3.font.size = Pt(9.5)
        p_dh3.font.color.rgb = CHARCOAL

    add_footer(slide6, 6)

    # =========================================================================
    # SLIDE 7: HOW TO ACCESS - ALUR 4 LANGKAH PENGGUNA
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(
        slide7,
        "HOW TO ACCESS — ALUR PENGGUNA (USER JOURNEY)",
        "Pengguna dapat mengakses platform dan memperoleh data melalui 4 tahapan alur swalayan yang terstruktur dan aman.",
        "Setiap tahapan dilengkapi validasi sistem otomatis untuk meminimalisasi kesalahan konfigurasi data."
    )

    steps = [
        (
            "LANGKAH 1",
            "Otentikasi & Profiling",
            "Login Single Sign-On (SSO)",
            [
                "Akses portal menggunakan peramban web modern.",
                "Login menggunakan akun SSO resmi Kementerian Keuangan.",
                "Sistem mengidentifikasi unit kerja dan menetapkan peran pengguna (RBAC) secara otomatis."
            ],
            NAVY_DARK
        ),
        (
            "LANGKAH 2",
            "Eksplorasi Katalog",
            "Pemilihan Entitas & Parameter",
            [
                "Pilih domain proses bisnis (Impor, Ekspor, Cukai, Fasilitas, Manifest, Perbendaharaan).",
                "Telusuri tabel dan pilih atribut kolom spesifik pada OpenMetadata Catalog.",
                "Tentukan rentang tanggal dan kriteria filter agregasi data."
            ],
            MCK_BLUE
        ),
        (
            "LANGKAH 3",
            "Justifikasi & PDP Review",
            "Penyusunan Formulir Dinas",
            [
                "Lengkapi latar belakang dinas, urgensi permohonan, dan estimasi volume baris.",
                "Sistem melakukan pra-pemeriksaan otomatis terhadap atribut sensitif (PII).",
                "Centang persetujuan pakta integritas kerahasiaan dan kepatuhan UU PDP."
            ],
            NAVY_MID
        ),
        (
            "LANGKAH 4",
            "Approval & Unduh Aman",
            "Verifikasi & Ekstraksi Data",
            [
                "Data Steward meneliti urgensi dan memberikan persetujuan digital (Digital Sign-Off).",
                "Engine mengeksekusi kueri teroptimasi pada CEISA DWH dengan masking dinamis.",
                "Unduh dataset dalam format Excel (.xlsx) atau ZIP terenkripsi melalui CMS Repositori."
            ],
            GREEN_ACC
        )
    ]

    for idx, (s_num, s_title, s_sub, s_points, s_color) in enumerate(steps):
        s_left = Inches(0.8 + idx * 3.0)
        s_top = Inches(1.9)
        s_w = Inches(2.8)
        s_h = Inches(4.3)

        card = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, s_left, s_top, s_w, s_h)
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        # Header step box
        h_box = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, s_left, s_top, s_w, Inches(0.95))
        h_box.fill.solid()
        h_box.fill.fore_color.rgb = s_color
        h_box.line.color.rgb = s_color

        tb_sh = slide7.shapes.add_textbox(s_left + Inches(0.15), s_top + Inches(0.08), s_w - Inches(0.3), Inches(0.8))
        tf_sh = tb_sh.text_frame
        tf_sh.word_wrap = True
        p_sh1 = tf_sh.paragraphs[0]
        p_sh1.text = s_num
        p_sh1.font.name = FONT_FAMILY
        p_sh1.font.size = Pt(8.5)
        p_sh1.font.bold = True
        p_sh1.font.color.rgb = RGBColor(226, 232, 240)

        p_sh2 = tf_sh.add_paragraph()
        p_sh2.text = s_title
        p_sh2.font.name = FONT_FAMILY
        p_sh2.font.size = Pt(11)
        p_sh2.font.bold = True
        p_sh2.font.color.rgb = WHITE

        # Body
        tb_sb = slide7.shapes.add_textbox(s_left + Inches(0.18), s_top + Inches(1.1), s_w - Inches(0.36), Inches(3.0))
        tf_sb = tb_sb.text_frame
        tf_sb.word_wrap = True
        tf_sb.margin_left = tf_sb.margin_right = tf_sb.margin_top = tf_sb.margin_bottom = 0

        p_subt = tf_sb.paragraphs[0]
        p_subt.text = s_sub
        p_subt.font.name = FONT_FAMILY
        p_subt.font.size = Pt(9.5)
        p_subt.font.bold = True
        p_subt.font.color.rgb = NAVY_DARK
        p_subt.space_after = Pt(8)

        for pt_idx, pt in enumerate(s_points):
            p_pt = tf_sb.add_paragraph()
            p_pt.text = f"•  {pt}"
            p_pt.font.name = FONT_FAMILY
            p_pt.font.size = Pt(9)
            p_pt.font.color.rgb = CHARCOAL
            p_pt.space_after = Pt(6)

    # Bottom Callout Summary
    bot_info7 = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.35), Inches(11.733), Inches(0.55))
    bot_info7.fill.solid()
    bot_info7.fill.fore_color.rgb = RGBColor(238, 242, 255)
    bot_info7.line.color.rgb = RGBColor(199, 210, 254)

    tb_bi7 = slide7.shapes.add_textbox(Inches(1.0), Inches(6.4), Inches(11.333), Inches(0.45))
    tf_bi7 = tb_bi7.text_frame
    tf_bi7.word_wrap = True
    p_bi7 = tf_bi7.paragraphs[0]
    p_bi7.text = "NOTIFIKASI REAL-TIME: Pemohon menerima notifikasi status langsung pada dashboard CMS di setiap perubahan tahapan review (Disetujui, Perlu Revisi, atau Ditolak)."
    p_bi7.font.name = FONT_FAMILY
    p_bi7.font.size = Pt(9.5)
    p_bi7.font.bold = True
    p_bi7.font.color.rgb = NAVY_MID

    add_footer(slide7, 7)

    # =========================================================================
    # SLIDE 8: HOW TO ACCESS - ROLE-BASED ACCESS CONTROL (RBAC)
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(
        slide8,
        "HOW TO ACCESS — TATA KELOLA PERAN (ROLE-BASED ACCESS CONTROL)",
        "Sistem membagi wewenang ke dalam 4 peran spesifik untuk menjamin prinsip 'Separation of Duties' dan keamanan data.",
        "Setiap pengguna memiliki ruang lingkup fitur yang dirancang khusus sesuai tanggung jawab organisasinya."
    )

    roles = [
        (
            "DATA REQUESTER",
            "Pemohon Unit Operasional / Kebijakan",
            "Staf / Pejabat di lingkungan DJBC atau Kementerian Keuangan yang membutuhkan data untuk analisis atau operasional.",
            [
                "Eksplorasi Katalog Enterprise OpenMetadata",
                "Penyusunan formulir permohonan data",
                "Pelacakan status tiket di CMS Repositori",
                "Pengunduhan berkas dataset yang telah disetujui"
            ],
            NAVY_MID
        ),
        (
            "DATA STEWARD",
            "Verifikator Kepatuhan & Tata Kelola",
            "Pejabat fungsional / pengelola data yang berwenang meneliti kelayakan permohonan dan kepatuhan perlindungan data.",
            [
                "Verifikasi justifikasi urgensi kedinasan",
                "Audit kepatuhan PDP dan peninjauan masking",
                "Penandatanganan digital persetujuan / penolakan",
                "Pemberian catatan revisi resmi kepada pemohon"
            ],
            MCK_BLUE
        ),
        (
            "DATA ANALYST",
            "Penyusun Model Data & Analitik",
            "Analis data dan riset makro yang memvalidasi logika agregasi dan relevansi data mart yang diminta.",
            [
                "Validasi sintaks SQL dan skema agregasi",
                "Pengujian konsistensi logika antar tabel",
                "Penyusunan data mart analitis terstandarisasi",
                "Optimalisasi dimensi dan metrik pelaporan"
            ],
            NAVY_DARK
        ),
        (
            "DATA ENGINEER",
            "Infrastruktur & Pipeline Data Warehouse",
            "Pengelola teknis basis data CEISA DWH yang memelihara kinerja server, kamus data, dan konektivitas database.",
            [
                "Sinkronisasi skema OpenMetadata dengan DWH fisik",
                "Monitoring performa eksekusi kueri dan latensi",
                "Pemeliharaan indeks tabel dan partisi data",
                "Audit keamanan konektivitas multi-dialek"
            ],
            SLATE
        )
    ]

    for idx, (r_name, r_sub, r_who, r_powers, r_color) in enumerate(roles):
        r_left = Inches(0.8 + idx * 3.0)
        r_top = Inches(1.9)
        r_w = Inches(2.8)
        r_h = Inches(4.8)

        card = slide8.shapes.add_shape(MSO_SHAPE.RECTANGLE, r_left, r_top, r_w, r_h)
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        # Header role
        rh_box = slide8.shapes.add_shape(MSO_SHAPE.RECTANGLE, r_left, r_top, r_w, Inches(0.85))
        rh_box.fill.solid()
        rh_box.fill.fore_color.rgb = r_color
        rh_box.line.color.rgb = r_color

        tb_rh = slide8.shapes.add_textbox(r_left + Inches(0.15), r_top + Inches(0.08), r_w - Inches(0.3), Inches(0.7))
        tf_rh = tb_rh.text_frame
        tf_rh.word_wrap = True
        p_rh1 = tf_rh.paragraphs[0]
        p_rh1.text = r_name
        p_rh1.font.name = FONT_FAMILY
        p_rh1.font.size = Pt(10)
        p_rh1.font.bold = True
        p_rh1.font.color.rgb = WHITE

        p_rh2 = tf_rh.add_paragraph()
        p_rh2.text = r_sub
        p_rh2.font.name = FONT_FAMILY
        p_rh2.font.size = Pt(8.5)
        p_rh2.font.color.rgb = RGBColor(226, 232, 240)

        # Body
        tb_rb = slide8.shapes.add_textbox(r_left + Inches(0.18), r_top + Inches(0.95), r_w - Inches(0.36), Inches(3.7))
        tf_rb = tb_rb.text_frame
        tf_rb.word_wrap = True
        tf_rb.margin_left = tf_rb.margin_right = tf_rb.margin_top = tf_rb.margin_bottom = 0

        p_rw = tf_rb.paragraphs[0]
        p_rw.text = "PROFIL PENGGUNA:"
        p_rw.font.name = FONT_FAMILY
        p_rw.font.size = Pt(8)
        p_rw.font.bold = True
        p_rw.font.color.rgb = SLATE

        p_rw2 = tf_rb.add_paragraph()
        p_rw2.text = r_who
        p_rw2.font.name = FONT_FAMILY
        p_rw2.font.size = Pt(8.5)
        p_rw2.font.color.rgb = CHARCOAL
        p_rw2.space_after = Pt(8)

        p_rp = tf_rb.add_paragraph()
        p_rp.text = "WEWENANG SISTEM:"
        p_rp.font.name = FONT_FAMILY
        p_rp.font.size = Pt(8)
        p_rp.font.bold = True
        p_rp.font.color.rgb = SLATE
        p_rp.space_after = Pt(2)

        for p_idx, power in enumerate(r_powers):
            p_p = tf_rb.add_paragraph()
            p_p.text = f"• {power}"
            p_p.font.name = FONT_FAMILY
            p_p.font.size = Pt(8.5)
            p_p.font.color.rgb = CHARCOAL
            p_p.space_after = Pt(3)

    add_footer(slide8, 8)

    # =========================================================================
    # SLIDE 9: HOW TO ACCESS - TITIK AKSES, PRASYARAT, & DUKUNGAN
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(
        slide9,
        "HOW TO ACCESS — TITIK AKSES, PRASYARAT, & DUKUNGAN TEKNIS",
        "Akses platform kini terbuka melalui portal resmi DataRequestStudio dengan panduan operasional terintegrasi.",
        "Informasi detail mengenai tautan sistem, prasyarat hak akses, dan kanal bantuan pengguna resmi."
    )

    access_boxes = [
        (
            "1. TITIK AKSES & TAUTAN PORTAL",
            "Alamat Portal Resmi & Lingkungan",
            [
                ("Portal Publik / Web Demo", "https://datalabs89.github.io/DataRequestStudio/"),
                ("Portal Intranet DJBC (Produksi)", "https://datarequest.customs.go.id (Jaringan Kemenkeu)"),
                ("Peramban yang Didukung", "Google Chrome, Microsoft Edge, Mozilla Firefox (Modern)"),
                ("Dukungan Antarmuka", "Responsif Desktop, Laptop Kantor, & Tablet Inspeksi")
            ],
            MCK_BLUE
        ),
        (
            "2. PRASYARAT & KELENGKAPAN AKSES",
            "Standar Administrasi Pengguna",
            [
                ("Otentikasi Akun", "Akun Kemenkeu ID / SSO Aktif"),
                ("Justifikasi Kedinasan", "Surat Perintah / Nota Dinas Kebutuhan Data"),
                ("Penetapan Peran (Role)", "Pengajuan hak akses ke Data Steward Unit bersangkutan"),
                ("Kepatuhan Tata Kelola", "Menandatangani Pakta Integritas Perlindungan Data Pribadi")
            ],
            NAVY_DARK
        ),
        (
            "3. LAYANAN BANTUAN & ASISTENSI",
            "Kanal Dukungan Operasional",
            [
                ("Unit Pengelola Utama", "Subdirektorat Pengelolaan Data & DWH, Dit. IKC DJBC"),
                ("Service Desk Khusus", "Email: data.steward@customs.go.id"),
                ("Helpdesk Tiket Internal", "Portal Service Desk Kemenkeu (Kategori: CEISA DWH)"),
                ("SLA Respon Bantuan", "Maksimal 2 jam kerja untuk kendala teknis akses")
            ],
            GREEN_ACC
        )
    ]

    for idx, (a_tag, a_sub, a_rows, a_color) in enumerate(access_boxes):
        c_left = Inches(0.8 + idx * 4.0)
        c_top = Inches(1.9)
        c_w = Inches(3.7)
        c_h = Inches(4.8)

        card = slide9.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_left, c_top, c_w, c_h)
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        # Header
        h_box = slide9.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_left, c_top, c_w, Inches(0.85))
        h_box.fill.solid()
        h_box.fill.fore_color.rgb = a_color
        h_box.line.color.rgb = a_color

        tb_ah = slide9.shapes.add_textbox(c_left + Inches(0.18), c_top + Inches(0.08), c_w - Inches(0.36), Inches(0.7))
        tf_ah = tb_ah.text_frame
        tf_ah.word_wrap = True
        p_ah1 = tf_ah.paragraphs[0]
        p_ah1.text = a_tag
        p_ah1.font.name = FONT_FAMILY
        p_ah1.font.size = Pt(9.5)
        p_ah1.font.bold = True
        p_ah1.font.color.rgb = WHITE

        p_ah2 = tf_ah.add_paragraph()
        p_ah2.text = a_sub
        p_ah2.font.name = FONT_FAMILY
        p_ah2.font.size = Pt(8.5)
        p_ah2.font.color.rgb = RGBColor(226, 232, 240)

        # Body key-value list
        tb_ab = slide9.shapes.add_textbox(c_left + Inches(0.18), c_top + Inches(1.0), c_w - Inches(0.36), Inches(3.6))
        tf_ab = tb_ab.text_frame
        tf_ab.word_wrap = True
        tf_ab.margin_left = tf_ab.margin_right = tf_ab.margin_top = tf_ab.margin_bottom = 0

        for r_idx, (label, val) in enumerate(a_rows):
            p_lbl = tf_ab.paragraphs[0] if r_idx == 0 else tf_ab.add_paragraph()
            p_lbl.text = label.upper()
            p_lbl.font.name = FONT_FAMILY
            p_lbl.font.size = Pt(8)
            p_lbl.font.bold = True
            p_lbl.font.color.rgb = a_color
            if r_idx > 0:
                p_lbl.space_before = Pt(8)

            p_val = tf_ab.add_paragraph()
            p_val.text = val
            p_val.font.name = FONT_FAMILY
            p_val.font.size = Pt(9.5)
            p_val.font.color.rgb = CHARCOAL
            p_val.space_before = Pt(1)

    add_footer(slide9, 9)

    # =========================================================================
    # SLIDE 10: KESIMPULAN & CALL TO ACTION (Next Steps)
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    add_header(
        slide10,
        "KESIMPULAN & LANGKAH TINDAK LANJUT",
        "Akselerasi pemanfaatan data kepabeanan yang bertanggung jawab, cepat, dan berintegritas dimulai dari DataRequestStudio.",
        "Tiga langkah strategis yang direkomendasikan bagi pimpinan unit kerja dan analis data."
    )

    steps_next = [
        (
            "1. MULAI EKSPLORASI MANDIRI",
            "Gunakan Katalog Terpadu",
            "Dorong staf analis di unit kerja Anda untuk mengeksplorasi kamus data dan skema di OpenMetadata Catalog untuk merumuskan kebutuhan data secara presisi sebelum mengajukan permohonan.",
            MCK_BLUE
        ),
        (
            "2. TERAPKAN MINIMISASI DATA",
            "Patuhi Prinsip Tata Kelola PDP",
            "Ajukan hanya atribut dan rentang waktu data yang mutlak diperlukan untuk kedinasan guna mempercepat proses verifikasi oleh Data Steward dan menjaga privasi wajib pajak.",
            GREEN_ACC
        ),
        (
            "3. SINERGI & INTEGRASI UNIT",
            "Onboarding & Umpan Balik Berkelanjutan",
            "Bagi unit kerja yang memiliki kebutuhan data mart khusus atau kueri analitik rutin, koordinasikan integrasi template permohonan bersama Subdirektorat Pengelolaan Data.",
            NAVY_DARK
        )
    ]

    for idx, (n_tag, n_head, n_desc, n_color) in enumerate(steps_next):
        c_left = Inches(0.8 + idx * 4.0)
        c_top = Inches(1.9)
        c_w = Inches(3.7)
        c_h = Inches(3.0)

        card = slide10.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_left, c_top, c_w, c_h)
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        # Stripe top
        st_top = slide10.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_left, c_top, c_w, Inches(0.1))
        st_top.fill.solid()
        st_top.fill.fore_color.rgb = n_color
        st_top.line.color.rgb = n_color

        tb_nb = slide10.shapes.add_textbox(c_left + Inches(0.2), c_top + Inches(0.25), c_w - Inches(0.4), c_h - Inches(0.4))
        tf_nb = tb_nb.text_frame
        tf_nb.word_wrap = True
        tf_nb.margin_left = tf_nb.margin_right = tf_nb.margin_top = tf_nb.margin_bottom = 0

        p_nt = tf_nb.paragraphs[0]
        p_nt.text = n_tag
        p_nt.font.name = FONT_FAMILY
        p_nt.font.size = Pt(9)
        p_nt.font.bold = True
        p_nt.font.color.rgb = n_color

        p_nh = tf_nb.add_paragraph()
        p_nh.text = n_head
        p_nh.font.name = FONT_FAMILY
        p_nh.font.size = Pt(12)
        p_nh.font.bold = True
        p_nh.font.color.rgb = NAVY_DARK
        p_nh.space_after = Pt(8)

        p_nd = tf_nb.add_paragraph()
        p_nd.text = n_desc
        p_nd.font.name = FONT_FAMILY
        p_nd.font.size = Pt(9.5)
        p_nd.font.color.rgb = CHARCOAL

    # Big Final Banner (Vision Callout)
    vis_box = slide10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(5.15), Inches(11.733), Inches(1.6))
    vis_box.fill.solid()
    vis_box.fill.fore_color.rgb = NAVY_DARK
    vis_box.line.color.rgb = NAVY_DARK

    tb_vb = slide10.shapes.add_textbox(Inches(1.1), Inches(5.35), Inches(11.133), Inches(1.2))
    tf_vb = tb_vb.text_frame
    tf_vb.word_wrap = True
    tf_vb.margin_left = tf_vb.margin_right = tf_vb.margin_top = tf_vb.margin_bottom = 0

    p_v1 = tf_vb.paragraphs[0]
    p_v1.text = "VISI DATAREQUESTSTUDIO: TATA KELOLA DATA PRIMA UNTUK KEPABEANAN & CUKAI"
    p_v1.font.name = FONT_FAMILY
    p_v1.font.size = Pt(10.5)
    p_v1.font.bold = True
    p_v1.font.color.rgb = MCK_BLUE

    p_v2 = tf_vb.add_paragraph()
    p_v2.text = "\"Data Tepat, Proses Cepat, Privasi Terjaga, Kebijakan Akurat.\""
    p_v2.font.name = FONT_FAMILY
    p_v2.font.size = Pt(18)
    p_v2.font.bold = True
    p_v2.font.color.rgb = WHITE
    p_v2.space_before = Pt(4)

    p_v3 = tf_vb.add_paragraph()
    p_v3.text = "Kunjungi portal resmi dan wujudkan tata kelola data berbasis bukti (evidence-based policy) di lingkungan Kementerian Keuangan."
    p_v3.font.name = FONT_FAMILY
    p_v3.font.size = Pt(10)
    p_v3.font.color.rgb = RGBColor(203, 213, 225)
    p_v3.space_before = Pt(4)

    add_footer(slide10, 10)

    # Save presentation
    prs.save(output_path)
    print(f"Presentation saved successfully to {os.path.abspath(output_path)}")
    return os.path.abspath(output_path)

if __name__ == "__main__":
    create_deck()
